#!/usr/bin/env python3
"""
Porsche Sanitizer Agent
Le uma planilha de vendas Porsche, aplica as regras do schema.md e grava um
novo .xlsx com as colunas *Sanitized* logo apos cada coluna de origem, mais
uma aba de relatorio de qualidade.

Uso:
    python porsche_sanitizer_agent.py entrada.xlsx [saida.xlsx] [--ref AAAA-MM-DD]

--ref define a data de referencia para datas futuras (padrao: hoje).
"""
import re
import sys
import calendar
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

INVALID = "INVALID"

# --------------------------------------------------------------------------- #
# Numeros por extenso
# --------------------------------------------------------------------------- #
UNITS = {w: i for i, w in enumerate(
    "zero one two three four five six seven eight nine ten eleven twelve thirteen "
    "fourteen fifteen sixteen seventeen eighteen nineteen".split())}
TENS = {"twenty": 20, "thirty": 30, "forty": 40, "fifty": 50,
        "sixty": 60, "seventy": 70, "eighty": 80, "ninety": 90}
SCALES = {"hundred": 100, "thousand": 1000, "million": 1_000_000}


def words_to_number(text):
    tokens = re.findall(r"[a-z]+", text.lower().replace("-", " "))
    tokens = [t for t in tokens if t != "and"]
    if not tokens or any(t not in UNITS and t not in TENS and t not in SCALES for t in tokens):
        return None
    total, current = 0, 0
    for t in tokens:
        if t in UNITS:
            current += UNITS[t]
        elif t in TENS:
            current += TENS[t]
        elif t == "hundred":
            current = (current or 1) * 100
        else:
            total += (current or 1) * SCALES[t]
            current = 0
    return total + current


# --------------------------------------------------------------------------- #
# Datas
# --------------------------------------------------------------------------- #
MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
MONTHS.update({m.lower(): i for i, m in enumerate(calendar.month_abbr) if m})
MONTHS["sept"] = 9


def _build_date(y, m, d):
    try:
        y, m, d = int(y), int(m), int(d)
        if y < 100:
            y += 2000
        return date(y, m, d).isoformat()
    except (ValueError, TypeError):
        return INVALID


def sanitize_date(value):
    if value is None or str(value).strip() == "":
        return INVALID
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    s = re.sub(r"\s+00:00:00$", "", s)

    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)  # YYYY-MM-DD
    if m:
        return _build_date(*m.groups())
    m = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{2}|\d{4})", s)  # MM/DD/YY(YY)
    if m:
        mo, d, y = m.groups()
        return _build_date(y, mo, d)
    m = re.fullmatch(r"([A-Za-z]+)\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})", s)  # Month DDth, YYYY
    if m and m.group(1).lower() in MONTHS:
        return _build_date(m.group(3), MONTHS[m.group(1).lower()], m.group(2))
    return INVALID


# Data de venda: alem de datas impossiveis, datas posteriores a data de
# referencia (padrao: dia da execucao) tambem viram INVALID.
REFERENCE_DATE = date.today()
DATE_ISSUES = {"impossivel": 0, "futura": 0}


def sanitize_sale_date(value):
    iso = sanitize_date(value)
    if iso == INVALID:
        DATE_ISSUES["impossivel"] += 1
        return INVALID
    if date.fromisoformat(iso) > REFERENCE_DATE:
        DATE_ISSUES["futura"] += 1
        return INVALID
    return iso


# --------------------------------------------------------------------------- #
# Modelos
# --------------------------------------------------------------------------- #
CANONICAL_MODELS = [
    "911 Carrera", "911 Carrera S", "911 Carrera GTS", "911 Turbo", "911 Turbo S",
    "911 GT3", "911 GT3 RS", "911 Dakar", "911 Targa 4", "911 Targa 4S",
    "718 Cayman", "718 Cayman S", "718 Cayman GT4 RS", "718 Boxster", "718 Boxster GTS",
    "718 Spyder RS", "Cayenne", "Cayenne S", "Cayenne Coupe", "Cayenne E-Hybrid",
    "Cayenne Turbo", "Cayenne Turbo GT", "Macan", "Macan S", "Macan T", "Macan GTS",
    "Macan Electric", "Panamera", "Panamera 4", "Panamera 4S", "Panamera Turbo",
    "Panamera Turbo S", "Panamera 4 E-Hybrid", "Taycan", "Taycan 4S", "Taycan GTS",
    "Taycan Turbo", "Taycan Turbo S", "Taycan Cross Turismo",
]
MODEL_LOOKUP = {re.sub(r"[\s\-_]+", " ", m.lower()): m for m in CANONICAL_MODELS}
UPPER_TOKENS = {"GT3", "GT4", "RS", "GTS", "S", "T", "4S", "E"}


def _smart_title(text):
    out = []
    for tok in text.split():
        parts = tok.split("-")
        fixed = []
        for p in parts:
            if p.upper() in UPPER_TOKENS or re.fullmatch(r"\d+[A-Za-z]?", p):
                fixed.append(p.upper())
            else:
                fixed.append(p.capitalize())
        out.append("-".join(fixed))
    return " ".join(out)


def sanitize_model(value):
    if value is None or str(value).strip() == "":
        return INVALID
    s = re.sub(r"\s+", " ", str(value)).strip()
    s = re.sub(r"^porsche\s+", "", s, flags=re.I)
    key = re.sub(r"[\s_]+", " ", s.lower()).replace("e hybrid", "e-hybrid")
    return MODEL_LOOKUP.get(key, _smart_title(s))


# --------------------------------------------------------------------------- #
# Ano do modelo
# --------------------------------------------------------------------------- #
def sanitize_year(value):
    if value is None or str(value).strip() == "":
        return INVALID
    s = str(value).strip().lower()
    s = re.sub(r"\.0$", "", s)
    year = None
    if re.fullmatch(r"\d{4}", s):
        year = int(s)
    elif re.fullmatch(r"\d{2}[\s\-/.]\d{2}", s):
        year = int(re.sub(r"\D", "", s))
    elif re.fullmatch(r"[a-z\s\-]+", s):
        words = s.replace("-", " ").split()
        # "twenty twenty four" = 20 + 24
        if len(words) >= 2 and words[0] in TENS and "thousand" not in words:
            hi = words_to_number(words[0])
            lo = words_to_number(" ".join(words[1:]))
            if hi is not None and lo is not None and lo < 100:
                year = hi * 100 + lo
        else:
            year = words_to_number(s)
    if year is None or not (1990 <= year <= 2035):
        return INVALID
    return str(year)


# --------------------------------------------------------------------------- #
# Numeros com separadores misturados (US e europeu)
# --------------------------------------------------------------------------- #
def parse_number(s):
    s = s.strip()
    if not re.fullmatch(r"\d[\d.,]*", s):
        return None
    has_dot, has_comma = "." in s, "," in s
    if has_dot and has_comma:
        dec = "." if s.rfind(".") > s.rfind(",") else ","
        thou = "," if dec == "." else "."
        s = s.replace(thou, "").replace(dec, ".")
    elif has_dot or has_comma:
        sep = "." if has_dot else ","
        groups = s.split(sep)
        # separador de milhar: todos os grupos apos o primeiro com 3 digitos
        if len(groups) > 1 and all(len(g) == 3 for g in groups[1:]) and len(groups[0]) <= 3:
            s = s.replace(sep, "")
        elif len(groups) == 2:
            s = groups[0] + "." + groups[1]
        else:
            return None
    try:
        return Decimal(s)
    except Exception:
        return None


# --------------------------------------------------------------------------- #
# Preco
# --------------------------------------------------------------------------- #
def sanitize_price(value):
    if value is None or str(value).strip() == "":
        return INVALID
    if isinstance(value, (int, float)):
        return str(Decimal(str(value)).quantize(Decimal("0.01"), ROUND_HALF_UP))
    s = str(value).lower()
    s = re.sub(r"usd|us\$|dollars?|\$", " ", s).strip()
    s = re.sub(r"\s+", " ", s)
    amount = None
    m = re.fullmatch(r"([\d.,]+)\s*k", s)
    if m:
        n = parse_number(m.group(1))
        amount = n * 1000 if n is not None else None
    elif re.fullmatch(r"[\d.,]+", s):
        amount = parse_number(s)
    elif re.fullmatch(r"[a-z\s\-]+", s):
        n = words_to_number(s)
        amount = Decimal(n) if n is not None else None
    if amount is None:
        return INVALID
    return str(amount.quantize(Decimal("0.01"), ROUND_HALF_UP))


# --------------------------------------------------------------------------- #
# Quilometragem
# --------------------------------------------------------------------------- #
KM_TO_MI = Decimal("0.621371")


def sanitize_mileage(value):
    if value is None or str(value).strip() == "":
        return INVALID
    if isinstance(value, (int, float)):
        return int(Decimal(str(value)).quantize(Decimal("1"), ROUND_HALF_UP))
    s = str(value).strip().lower()
    is_km = bool(re.search(r"\bkm\b|kilomet", s))
    if s in {"new", "new car", "zero", "zero miles", "brand new"}:
        return 0
    s = re.sub(r"kilometers?|kilometres?|\bkm\b|miles?|\bmi\b\.?|mi\.?$|:", " ", s)
    s = re.sub(r"(\d)mi\.?", r"\1", s)
    s = re.sub(r"\s+", " ", s).strip()
    n = None
    if re.fullmatch(r"[\d.,]+", s):
        n = parse_number(s)
    elif re.fullmatch(r"[a-z\s\-]+", s):
        w = words_to_number(s)
        n = Decimal(w) if w is not None else None
    if n is None:
        return INVALID
    if is_km:
        n = n * KM_TO_MI
    return int(n.quantize(Decimal("1"), ROUND_HALF_UP))


# --------------------------------------------------------------------------- #
# Pagamento
# --------------------------------------------------------------------------- #
PAYMENT_RULES = [
    (r"crypto|bitcoin|btc", "Crypto Payment"),
    (r"\bach\b", "ACH Payment"),
    (r"debit", "Debit Card"),
    (r"credit", "Credit Card"),
    (r"wire", "Wire Transfer"),
    (r"bank ?transfer|transfer", "Bank Transfer"),
    (r"financ|loan", "Financing"),
    (r"leas", "Lease"),
    (r"cash", "Cash"),
]


def sanitize_payment(value):
    if value is None or str(value).strip() == "":
        return INVALID
    raw = str(value).strip()
    s = re.sub(r"([a-z])([A-Z])", r"\1 \2", raw)          # CreditCard -> Credit Card
    s = re.sub(r"[\-_]+", " ", s).lower()
    s = re.sub(r"\s+", " ", s).strip()
    for pattern, label in PAYMENT_RULES:
        if re.search(pattern, s):
            return label
    return s.title()


# --------------------------------------------------------------------------- #
# Cidade
# --------------------------------------------------------------------------- #
CITY_EXCEPTIONS = {"st louis": "St. Louis", "st. louis": "St. Louis",
                   "saint louis": "St. Louis", "st paul": "St. Paul", "st. paul": "St. Paul",
                   "st petersburg": "St. Petersburg", "st. petersburg": "St. Petersburg"}


def sanitize_city(value):
    if value is None or str(value).strip() == "":
        return INVALID
    s = re.sub(r"\s+", " ", str(value)).strip()
    if s.lower() in CITY_EXCEPTIONS:
        return CITY_EXCEPTIONS[s.lower()]
    return " ".join(w[:1].upper() + w[1:].lower() for w in s.split(" "))


# --------------------------------------------------------------------------- #
# Estado
# --------------------------------------------------------------------------- #
STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA",
    "colorado": "CO", "connecticut": "CT", "delaware": "DE", "florida": "FL", "georgia": "GA",
    "hawaii": "HI", "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA",
    "kansas": "KS", "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN", "mississippi": "MS",
    "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV", "new hampshire": "NH",
    "new jersey": "NJ", "new mexico": "NM", "new york": "NY", "north carolina": "NC",
    "north dakota": "ND", "ohio": "OH", "oklahoma": "OK", "oregon": "OR", "pennsylvania": "PA",
    "rhode island": "RI", "south carolina": "SC", "south dakota": "SD", "tennessee": "TN",
    "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY", "district of columbia": "DC",
}
CODES = set(STATES.values())


def sanitize_state(value):
    if value is None or str(value).strip() == "":
        return INVALID
    s = re.sub(r"[.\s]+", " ", str(value)).strip().lower()
    if s.upper().replace(" ", "") in CODES and len(s.replace(" ", "")) == 2:
        return s.upper().replace(" ", "")
    return STATES.get(s, INVALID)


# --------------------------------------------------------------------------- #
# Status de entrega
# --------------------------------------------------------------------------- #
DELIVERY_LABELS = ["Delivered", "Pending", "In Transit", "Cancelled", "Awaiting Delivery",
                   "Awaiting Pickup", "Pending Approval", "Pending Review", "Shipped",
                   "Awaiting Review"]
DELIVERY_LOOKUP = {l.lower(): l for l in DELIVERY_LABELS}
DELIVERY_LOOKUP.update({"deliverd": "Delivered", "delievered": "Delivered",
                        "canceled": "Cancelled", "intransit": "In Transit"})


def sanitize_delivery(value):
    if value is None or str(value).strip() == "":
        return INVALID
    s = re.sub(r"[^a-z\s\-_]", "", str(value).lower())
    s = re.sub(r"[\-_\s]+", " ", s).strip()
    if s in DELIVERY_LOOKUP:
        return DELIVERY_LOOKUP[s]
    if s.replace(" ", "") in DELIVERY_LOOKUP:
        return DELIVERY_LOOKUP[s.replace(" ", "")]
    return s.title() if s else INVALID


# --------------------------------------------------------------------------- #
# Pipeline
# --------------------------------------------------------------------------- #
RULES = [  # (coluna de origem, coluna de saida, funcao)
    ("sale_date", "SaleDateSanitized", sanitize_sale_date),
    ("porsche_model", "PorscheModelSanitized", sanitize_model),
    ("model_year", "ModelYearSanitized", sanitize_year),
    ("sale_price", "SalesPriceSanitized", sanitize_price),
    ("vehicle_mileage", "VehicleMileageSanitized", sanitize_mileage),
    ("payment_method", "PayMethodSanitized", sanitize_payment),
    ("city", "CitySanitized", sanitize_city),
    ("state", "StateSanitized", sanitize_state),
    ("delivery_status", "DeliveryStatusSanitized", sanitize_delivery),
]
REQUIRED = ["sale_id", "sale_date", "customer_name", "porsche_model", "model_year",
            "sale_price", "vehicle_mileage", "payment_method", "city", "state",
            "salesperson", "delivery_status"]


def run(input_path, output_path):
    DATE_ISSUES["impossivel"] = 0
    DATE_ISSUES["futura"] = 0
    wb = load_workbook(input_path, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    data = [r for r in rows[1:] if any(v not in (None, "") for v in r)]

    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise SystemExit(f"Colunas obrigatorias ausentes: {missing}")

    rule_by_src = {src: (out, fn) for src, out, fn in RULES}
    out_header = []
    for h in header:
        out_header.append(h)
        if h in rule_by_src:
            out_header.append(rule_by_src[h][0])

    out_rows = []
    for r in data:
        rec = dict(zip(header, r))
        new = []
        for h in header:
            v = rec.get(h)
            new.append(v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, datetime) else v)
            if h in rule_by_src:
                new.append(rule_by_src[h][1](v))
        out_rows.append(new)

    checks = quality_checks(header, out_header, out_rows)
    write_output(output_path, ws.title, out_header, out_rows, checks)
    return checks, out_header, out_rows


def quality_checks(header, out_header, out_rows):
    idx = {h: i for i, h in enumerate(out_header)}
    col = lambda name: [row[idx[name]] for row in out_rows]
    results = []

    def add(name, ok, detail=""):
        results.append((name, "OK" if ok else "FALHOU", detail))

    add("Nenhuma coluna original removida", all(h in out_header for h in header))
    add("Colunas sanitizadas logo apos a origem",
        all(out_header[idx[src] + 1] == out for src, out, _ in RULES))
    add("Datas em ISO ou INVALID",
        all(v == INVALID or re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) for v in col("SaleDateSanitized")))
    add("Estados com codigo USPS ou INVALID",
        all(v == INVALID or v in CODES for v in col("StateSanitized")))
    add("Quilometragem inteira",
        all(isinstance(v, int) or v == INVALID for v in col("VehicleMileageSanitized")))
    add("Preco numerico com duas casas",
        all(v == INVALID or re.fullmatch(r"\d+\.\d{2}", v) for v in col("SalesPriceSanitized")))
    add("Anos com 4 digitos ou INVALID",
        all(v == INVALID or re.fullmatch(r"\d{4}", v) for v in col("ModelYearSanitized")))
    add("Pagamento em rotulo controlado",
        all(v in {l for _, l in PAYMENT_RULES} for v in col("PayMethodSanitized")),
        "valores fora da lista sao title-case conforme o schema")
    add("Entrega em rotulo controlado",
        all(v in DELIVERY_LABELS for v in col("DeliveryStatusSanitized")))
    add("Nenhum campo sanitizado em branco",
        all(row[idx[out]] not in (None, "") for row in out_rows for _, out, _ in RULES))

    for _, out, _ in RULES:
        n = sum(1 for v in col(out) if v == INVALID)
        results.append((f"Contagem de INVALID em {out}", "INFO", str(n)))
    results.append(("Datas de venda impossiveis ou ilegiveis", "INFO",
                    str(DATE_ISSUES["impossivel"])))
    results.append(("Datas de venda futuras", "INFO",
                    f'{DATE_ISSUES["futura"]} (posteriores a {REFERENCE_DATE.isoformat()})'))
    return results


def write_output(path, sheet_name, out_header, out_rows, checks):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(out_header)
    head_font = Font(bold=True, color="FFFFFF")
    raw_fill = PatternFill("solid", fgColor="404040")
    san_fill = PatternFill("solid", fgColor="1F6E43")
    inv_fill = PatternFill("solid", fgColor="F8D7DA")
    for c, h in enumerate(out_header, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = san_fill if h.endswith("Sanitized") else raw_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in out_rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            if cell.value == INVALID:
                cell.fill = inv_fill
            if out_header[cell.column - 1] in ("SalesPriceSanitized", "ModelYearSanitized"):
                cell.number_format = "@"
    for c, h in enumerate(out_header, 1):
        width = max([len(str(h))] + [len(str(row[c - 1])) for row in out_rows if row[c - 1] is not None])
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = min(width + 2, 40)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    rep = wb.create_sheet("Relatorio_Qualidade")
    rep.append(["Verificacao", "Resultado", "Detalhe"])
    for c in range(1, 4):
        rep.cell(row=1, column=c).font = head_font
        rep.cell(row=1, column=c).fill = san_fill
    for chk in checks:
        rep.append(list(chk))
    rep.column_dimensions["A"].width = 48
    rep.column_dimensions["B"].width = 12
    rep.column_dimensions["C"].width = 50
    wb.save(path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    args = sys.argv[1:]
    if "--ref" in args:
        i = args.index("--ref")
        REFERENCE_DATE = date.fromisoformat(args[i + 1])
        del args[i:i + 2]
    src = Path(args[0])
    dst = Path(args[1]) if len(args) > 1 else src.with_name(src.stem + "_sanitized.xlsx")
    checks, _, _ = run(src, dst)
    for name, status, detail in checks:
        print(f"[{status}] {name} {detail}")
    print(f"\nArquivo gerado: {dst}")